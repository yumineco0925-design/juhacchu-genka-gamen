"""サンプルデータ assets/data.js から、工事別の原価表 docs/02_genka_sample.xlsx を作る。

集計は data.js の表だけを使ってここで行う。言葉の定義は次のとおり。
  予算       当初の予算に、承認済の予算改定の差（改定後 − 改定前）を足した額
  有効発注額 発注金額に、その発注を元にした承認済の減額発注を足した額
  発注済     承認済と検収済の発注の有効発注額と、打切の発注の検収累計の和。労務費は計上済の労務費の和
  発生原価   有効な発注検収の累計。労務費は計上済の労務費
  残         予算 − 発注済
  消化率     発注済 ÷ 予算。予算が 0 のときは無し
  粗利（発生原価ベース） 受注金額 − 発生原価の合計。受注金額が無ければ無し

使い方: python make_xlsx.py
"""
import json
import os
import sys
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.normpath(os.path.join(HERE, '..'))
DATA = os.path.join(SITE, 'assets', 'data.js')
OUT = os.path.join(SITE, 'docs', '02_genka_sample.xlsx')
PREFIX = 'window.SAMPLE = '
HEADER = '架空の会社の資料です（自主制作）'
TITLE = '工事別の原価表'
SHEET = '原価表'
CATS = ['材料費', '外注費', '労務費', '経費']

# 色は藍の1系統と無彩色だけ
INDIGO = 'FF1D4478'
INDIGO_TINT = 'FFE8EDF4'
GRAY_BG = 'FFF3F3F3'
GRAY_LINE = 'FFBFBFBF'
INK_2 = 'FF4D4D4D'
FONT = '游ゴシック'
NUM_FMT = '#,##0'

HEAD_ROWS = (4, 5)


def load(path=DATA):
    """data.js の「window.SAMPLE = 」の後を JSON として読む。"""
    with open(path, encoding='utf-8') as f:
        text = f.read()
    body = text[text.index(PREFIX) + len(PREFIX):].rstrip()
    if body.endswith(';'):
        body = body[:-1]
    return json.loads(body)


def aggregate(S):
    """工事ごとに、費目別の 予算・発注済・発生原価・残・消化率 と合計、粗利（発生原価ベース）を出す。"""
    orders = S['orders']

    def eff(o):
        return o['amount'] + sum(r['amount'] for r in orders
                                 if r['related'] == o['no'] and r['amount'] < 0 and r['status'] == '承認済')

    cum = {}
    for i in S['inspections']:
        if i['status'] == '有効':
            cum[i['order']] = cum.get(i['order'], 0) + i['amount']

    result = []
    for k in S['kouji']:
        rows = []
        for cat in CATS:
            budget = (k['budget0'] or {}).get(cat, 0) + sum(
                r['after'] - r['before'] for r in S['revisions']
                if r['kouji'] == k['no'] and r['cat'] == cat and r['status'] == '承認済')
            if cat == '労務費':
                committed = sum(l['amount'] for l in S['labor'] if l['kouji'] == k['no'] and l['recordedOn'])
                incurred = committed
            else:
                # 減額発注は元の発注の有効発注額に入っているので、ここでは数えない
                mine = [o for o in orders if o['kouji'] == k['no'] and o['cat'] == cat and o['amount'] >= 0]
                committed = (sum(eff(o) for o in mine if o['status'] in ('承認済', '検収済'))
                             + sum(cum.get(o['no'], 0) for o in mine if o['status'] == '打切'))
                incurred = sum(cum.get(o['no'], 0) for o in mine)
            rows.append({'cat': cat, 'budget': budget, 'committed': committed, 'incurred': incurred,
                         'remain': budget - committed, 'rate': committed / budget if budget else None})
        total = {'cat': '合計'}
        for f in ('budget', 'committed', 'incurred', 'remain'):
            total[f] = sum(r[f] for r in rows)
        total['rate'] = total['committed'] / total['budget'] if total['budget'] else None
        contract = k['contract']
        result.append({'no': k['no'], 'name': k['name'], 'status': k['status'], 'contract': contract,
                       'rows': rows, 'total': total,
                       'profit': None if contract is None else contract - total['incurred']})
    return result


def columns():
    """列の見出し（上の段、下の段）の並び。"""
    cols = [('工事番号', ''), ('工事名', ''), ('状態', ''), ('受注金額（円）', '')]
    for label in CATS + ['合計']:
        cols += [(f'{label}（円）', '予算'), (f'{label}（円）', '発注済')]
    cols += [('残（円）', ''), ('粗利（発生原価ベース、円）', '')]
    return cols


def table_rows(result):
    """見出しの下に書く行（工事の行と合計の行）を値の並びで返す。"""
    rows = []
    for a in result:
        line = [a['no'], a['name'], a['status'], a['contract']]
        for r in a['rows']:
            line += [r['budget'], r['committed']]
        line += [a['total']['budget'], a['total']['committed'], a['total']['remain'], a['profit']]
        rows.append(line)
    sums = ['合計', None, None]
    for c in range(3, len(rows[0])):
        sums.append(sum(line[c] for line in rows if line[c] is not None))
    rows.append(sums)
    return rows


def text_width(s):
    """全角を2、半角を1と数えた表示の幅。"""
    return sum(2 if ord(ch) > 0xFF else 1 for ch in str(s))


def build(S):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    cols = columns()
    ncol = len(cols)
    body = table_rows(aggregate(S))

    thin = Side(style='thin', color=GRAY_LINE)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor=INDIGO_TINT)
    sum_fill = PatternFill('solid', fgColor=GRAY_BG)
    f_body = Font(name=FONT, size=10)
    f_head = Font(name=FONT, size=10, color=INDIGO)
    f_sum = Font(name=FONT, size=10)

    # 1行目は断り、2行目は表題と基準日
    ws['A1'] = HEADER
    ws['A1'].font = Font(name=FONT, size=9, color=INK_2)
    ws['A2'] = TITLE
    ws['A2'].font = Font(name=FONT, size=12, bold=True, color=INDIGO)
    right = ws.cell(2, ncol, f'基準日 {S["settings"]["baseDate"]}　金額は税抜')
    right.font = Font(name=FONT, size=10, color=INK_2)
    right.alignment = Alignment(horizontal='right')

    # 見出しは2段。費目と合計は上の段で2列をまとめ、下の段に 予算 と 発注済
    h1, h2 = HEAD_ROWS
    for c, (top, sub) in enumerate(cols, start=1):
        for r, text in ((h1, top), (h2, sub)):
            cell = ws.cell(r, c, text or None)
            cell.font = f_head
            cell.fill = head_fill
            cell.border = box
            if c <= 3:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            elif sub and r == h1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
        if not sub:
            ws.cell(h2, c).value = None
            ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
    for c in range(5, ncol - 1, 2):
        ws.merge_cells(start_row=h1, start_column=c, end_row=h1, end_column=c + 1)

    # 工事の行と合計の行。金額は3桁区切りで右揃え
    start = h2 + 1
    for i, line in enumerate(body):
        r = start + i
        is_sum = i == len(body) - 1
        for c, v in enumerate(line, start=1):
            cell = ws.cell(r, c, v)
            cell.border = box
            cell.font = f_sum if is_sum else f_body
            if is_sum:
                cell.fill = sum_fill
            if c >= 4:
                cell.number_format = NUM_FMT
                cell.alignment = Alignment(horizontal='right')
    ws.merge_cells(start_row=start + len(body) - 1, start_column=1,
                   end_row=start + len(body) - 1, end_column=3)

    # 列の幅は中身の最長に合わせる。金額は3桁区切りにした桁数で数える
    for c, (top, sub) in enumerate(cols, start=1):
        width = text_width(sub) if sub else text_width(top)
        if sub:
            width = max(width, (text_width(top) + 1) // 2)
        for line in body:
            v = line[c - 1]
            if v is not None:
                width = max(width, text_width(f'{v:,}' if isinstance(v, int) else v))
        ws.column_dimensions[get_column_letter(c)].width = width + 3
    ws.row_dimensions[h1].height = 20
    ws.row_dimensions[h2].height = 20

    # 見出しの下と工事名の右で固定する。印刷は横向きで幅を1枚に収める
    ws.freeze_panes = ws.cell(start, 3)
    ws.print_title_rows = f'{h1}:{h2}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 作成者などの文書の情報は空にする
    p = wb.properties
    p.creator = ''
    p.lastModifiedBy = ''
    p.title = None
    p.subject = None
    p.description = None
    p.keywords = None
    return wb


def check_props(path):
    """保存した xlsx の core.xml と app.xml で、作成者・最終更新者・会社・管理者が空かを確かめる。"""
    wanted = {'docProps/core.xml': ('creator', 'lastModifiedBy'),
              'docProps/app.xml': ('Company', 'Manager')}
    bad = []
    with zipfile.ZipFile(path) as z:
        for part, names in wanted.items():
            root = ET.fromstring(z.read(part))
            for el in root.iter():
                tag = el.tag.split('}')[-1]
                if tag in names and (el.text or '').strip():
                    bad.append(f'{part} {tag}')
    return bad


def main():
    S = load()
    build(S).save(OUT)
    bad = check_props(OUT)
    if bad:
        print('文書の情報が空になっていない: ' + '、'.join(bad))
        return 1
    ws = load_workbook(OUT)[SHEET]
    if ws['A1'].value != HEADER:
        print('1行目の断りを読み戻せない')
        return 1
    print('wrote', os.path.relpath(OUT, SITE), f'({ws.max_row} 行 x {ws.max_column} 列)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
